"""Génère le classeur Excel de publipostage « Courrier augmentation tarifaire ».

Usage : python generer_courrier.py  ->  Courrier_augmentation_tarifaire.xlsx
"""
import math

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

SORTIE = "Courrier_augmentation_tarifaire.xlsx"
FONT = "Arial"

BLEU_SAISIE = "0000FF"
JAUNE = PatternFill("solid", fgColor="FFF2CC")
BLEU_FONCE = "1F3864"
ENTETE = PatternFill("solid", fgColor=BLEU_FONCE)
GRIS_CLAIR = PatternFill("solid", fgColor="F2F2F2")
fin = Side(style="thin", color="BFBFBF")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
EUR = '#,##0" €";-#,##0" €";"—"'
PCT = '0.0%;-0.0%;"—"'


def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


wb = Workbook()

# ---------------------------------------------------------------- Paramètres
wp = wb.active
wp.title = "Paramètres"
wp.column_dimensions["A"].width = 34
wp.column_dimensions["B"].width = 110
wp["A1"] = "PARAMÈTRES DU COURRIER"
wp["A1"].font = f(14, True, BLEU_FONCE)
wp["A2"] = ("Cellules sur fond jaune / texte bleu = à personnaliser. "
            "Balises utilisables dans les textes : {cabinet} {societe} {date_effet} {hausse}")
wp["A2"].font = f(9, italic=True, color="7F7F7F")

params = [
    # (nom défini, libellé, valeur, format, commentaire)
    ("__titre", "Coordonnées du cabinet", None, None, None),
    ("Cab_Nom", "Nom du cabinet", "ALTHO EXPERTISE", None, None),
    ("Cab_Adr1", "Adresse", "141 rue Charles de Gaulle", None, None),
    ("Cab_CPVille", "Code postal et ville", "95130 LE PLESSIS-BOUCHARD", None, None),
    ("Cab_Ville", "Ville (lieu de signature)", "Le Plessis-Bouchard", None, None),
    ("Cab_Contact", "Téléphone / e-mail", "Tél. 01 34 14 53 45  ·  altho@altho-experts.com  ·  altho-experts.com", None, None),
    ("Cab_Mentions", "Mentions légales (pied de page)",
     "ALTHO EXPERTISE · SAS au capital de 10 000 € · RCS Pontoise 919 706 325 · TVA FR62919706325 · "
     "Société d'expertise comptable inscrite au Tableau de l'Ordre des experts-comptables de Paris Île-de-France",
     None, "Source : registre du commerce (Pappers). Vérifiez l'intitulé exact de l'inscription à l'Ordre."),
    ("Signataire", "Nom du signataire", "Prénom NOM", None, None),
    ("Signataire_Titre", "Fonction du signataire", "Expert-comptable associé", None, None),
    ("__titre", "Dates", None, None, None),
    ("Date_Courrier", "Date du courrier", "=TODAY()", "dd/mm/yyyy",
     "Par défaut : date du jour. Vous pouvez saisir une date fixe (ex. 01/12/2026)."),
    ("Date_Effet", "Date d'effet des nouveaux tarifs", "__DATE_EFFET", "dd/mm/yyyy", None),
    ("__titre", "Taux d'augmentation par défaut (par mission)", None, None, None),
    ("Taux_Compta", "Comptabilité", 0.035, "0.0%",
     "Taux appliqué à tous les clients, sauf taux personnalisé saisi dans l'onglet Clients."),
    ("Taux_Paye", "Paye / social", 0.04, "0.0%", None),
    ("Taux_Jur", "Juridique", 0.03, "0.0%", None),
    ("Taux_PA", "Plateforme agréée (facturation électronique)", 0.0, "0.0%", None),
    ("Pas_Arrondi", "Arrondi des nouveaux honoraires (en €)", 1, "0",
     "1 = à l'euro près, 5 = aux 5 € supérieurs, 10 = aux 10 € supérieurs…"),
    ("Nb_Mensualites", "Nombre d'échéances par an (mensualisation)", 12, "0",
     "Utilisé pour afficher le montant de la mensualité HT sous le tableau. 0 = ligne masquée."),
    ("Taux_TVA", "Taux de TVA", 0.2, "0.0%", None),
    ("__titre", "Impression en masse", None, None, None),
    ("Filtre_Masse", "Clients à inclure", "Courrier postal uniquement", None,
     "Courrier postal uniquement = clients « À envoyer = Oui » avec mode Courrier ou Les deux.\n"
     "Tous les clients à envoyer = tous les clients « À envoyer = Oui », quel que soit le mode."),
    ("__titre", "Libellés des missions (tels qu'imprimés dans le courrier)", None, None, None),
    ("Lib_Compta", "Mission 1", "Mission comptable (tenue, révision, bilan, liasse fiscale)", None, None),
    ("Lib_Paye", "Mission 2", "Mission sociale (bulletins de paie, DSN, déclarations sociales)", None, None),
    ("Lib_Jur", "Mission 3", "Mission juridique (approbation des comptes, secrétariat juridique)", None, None),
    ("Lib_PA", "Mission 4", "Plateforme agréée – facturation électronique", None, None),
    ("__titre", "Textes du courrier (modifiables — laissez vide pour supprimer un paragraphe)", None, None, None),
    ("Txt_Objet", "Objet", "Objet : révision de nos honoraires au {date_effet}", None, None),
    ("Txt_P1", "Paragraphe 1",
     "Nous tenons tout d'abord à vous remercier pour la confiance que vous accordez à {cabinet}. "
     "Comme chaque année, et conformément aux dispositions de votre lettre de mission, nous procédons "
     "à la révision de nos honoraires.", None, None),
    ("Txt_P2", "Paragraphe 2",
     "Cette révision, que nous avons souhaitée mesurée, tient compte de l'évolution de nos charges "
     "(rémunération et formation continue de nos collaborateurs, logiciels, sécurité informatique) "
     "ainsi que du renforcement constant des obligations comptables, fiscales, sociales et juridiques "
     "qui s'imposent à votre entreprise.", None, None),
    ("Txt_P3", "Paragraphe 3 (facturation électronique)",
     "L'année à venir est également marquée par la généralisation de la facturation électronique : "
     "depuis le 1er septembre 2026, toutes les entreprises doivent être en mesure de recevoir leurs "
     "factures électroniques via une plateforme agréée, et l'obligation d'émission s'étendra aux PME "
     "et micro-entreprises au 1er septembre 2027. Notre cabinet vous accompagne dans cette transition "
     "grâce à une solution de plateforme agréée directement intégrée à votre dossier.", None, None),
    ("Txt_P4", "Paragraphe 4 (avant le tableau)",
     "À compter du {date_effet}, vos honoraires annuels hors taxes évolueront ainsi "
     "(évolution globale : {hausse}) :", None, None),
    ("Txt_P5", "Paragraphe 5 (après le tableau)",
     "Les autres conditions de votre lettre de mission demeurent inchangées. Les nouveaux montants "
     "seront appliqués automatiquement sur vos prochaines factures et échéances de prélèvement.",
     None, None),
    ("Txt_P6", "Paragraphe 6",
     "Votre interlocuteur habituel reste à votre entière disposition pour répondre à vos questions "
     "et faire le point sur l'évolution de vos besoins.", None, None),
    ("Txt_Politesse", "Formule de politesse",
     "Nous vous prions d'agréer, {civilite}, l'expression de nos salutations distinguées.", None,
     "{civilite} est remplacé par Madame / Monsieur / Madame, Monsieur selon la civilité du client."),
]

r = 4
nom_vers_cellule = {}
for nom, lib, val, fmt, com in params:
    if nom == "__titre":
        r += 1
        wp.cell(r, 1, lib).font = f(11, True, "FFFFFF")
        wp.cell(r, 1).fill = ENTETE
        wp.cell(r, 2).fill = ENTETE
        r += 1
        continue
    wp.cell(r, 1, lib).font = f(10)
    c = wp.cell(r, 2)
    if val == "__DATE_EFFET":
        c.value = "=DATE(YEAR(Date_Courrier)+1,1,1)"
        c.comment = Comment("Par défaut : 1er janvier de l'année suivante. Saisissez une autre date si besoin.",
                            "Modèle")
    else:
        c.value = val
    c.font = f(10, color=BLEU_SAISIE)
    c.fill = JAUNE
    c.border = BORD
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    if fmt:
        c.number_format = fmt
    if com:
        c.comment = Comment(com, "Modèle")
    if isinstance(val, str) and len(val) > 100:
        wp.row_dimensions[r].height = 13 * math.ceil(len(val) / 105) + 3
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Paramètres'!$B${r}")
    nom_vers_cellule[nom] = f"'Paramètres'!$B${r}"
    r += 1

# Liste des mois (pour écrire les dates en toutes lettres, indépendamment de la langue d'Excel)
r += 2
wp.cell(r, 1, "Mois (usage interne)").font = f(9, italic=True, color="7F7F7F")
mois = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre"]
for i, m in enumerate(mois):
    wp.cell(r + i, 2, m).font = f(9, color="7F7F7F")
wb.defined_names["Liste_Mois"] = DefinedName("Liste_Mois", attr_text=f"'Paramètres'!$B${r}:$B${r + 11}")
wp.freeze_panes = "A4"
dv_filtre = DataValidation(type="list", formula1='"Courrier postal uniquement,Tous les clients à envoyer"')
wp.add_data_validation(dv_filtre)
dv_filtre.add(nom_vers_cellule["Filtre_Masse"].split("!")[1].replace("$", ""))

# ---------------------------------------------------------------- Clients
wc = wb.create_sheet("Clients")
colonnes = [
    # (entête, largeur, type)  type: saisie / calc
    ("Code client", 12, "saisie"),         # A
    ("Civilité", 9, "saisie"),             # B
    ("Prénom", 13, "saisie"),              # C
    ("Nom", 16, "saisie"),                 # D
    ("Société", 28, "saisie"),             # E
    ("Adresse 1", 28, "saisie"),           # F
    ("Adresse 2", 18, "saisie"),           # G
    ("Code postal", 10, "saisie"),         # H
    ("Ville", 16, "saisie"),               # I
    ("Email", 26, "saisie"),               # J
    ("Compta actuel HT", 13, "saisie"),    # K
    ("Paye actuel HT", 13, "saisie"),      # L
    ("Juridique actuel HT", 13, "saisie"),  # M
    ("PA actuel HT", 13, "saisie"),        # N
    ("Taux perso (optionnel)", 12, "saisie"),  # O
    ("Compta nouveau HT", 13, "calc"),     # P
    ("Paye nouveau HT", 13, "calc"),       # Q
    ("Juridique nouveau HT", 13, "calc"),  # R
    ("PA nouveau HT", 13, "calc"),         # S
    ("Total actuel HT", 13, "calc"),       # T
    ("Total nouveau HT", 13, "calc"),      # U
    ("Écart HT", 11, "calc"),              # V
    ("Écart %", 9, "calc"),                # W
    ("À envoyer", 10, "saisie"),           # X
    ("Mode d'envoi", 12, "saisie"),        # Y
    ("N° impression en masse", 12, "calc"),  # Z
]
for i, (titre, larg, typ) in enumerate(colonnes, start=1):
    c = wc.cell(1, i, titre)
    c.font = f(10, True, "FFFFFF")
    c.fill = ENTETE if typ == "saisie" else PatternFill("solid", fgColor="548235")
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    c.border = BORD
    wc.column_dimensions[get_column_letter(i)].width = larg
wc.row_dimensions[1].height = 32
wc.freeze_panes = "F2"

exemples = [
    ["C001", "M.", "Pierre", "DURAND", "SARL DURAND BÂTIMENT", "5 avenue des Tilleuls", "", "69100",
     "Villeurbanne", "p.durand@exemple.fr", 3600, 1800, 600, 240, None, "Oui", "Courrier"],
    ["C002", "Mme", "Sophie", "LEROY", "SAS LEROY CONSEIL", "Parc d'activités du Moulin", "Bât. B – 3 allée Verte",
     "38000", "Grenoble", "s.leroy@exemple.fr", 5400, 0, 900, 0, 0.025, "Oui", "Email"],
    ["C003", "", "", "", "EURL LES JARDINS DE CLAIRE", "18 chemin des Vignes", "", "69400",
     "Villefranche-sur-Saône", "contact@jardins-claire.fr", 2400, 1200, 0, 180, None, "Oui", "Les deux"],
]
NB_LIGNES = 500
TAUX = {"P": ("K", "Taux_Compta"), "Q": ("L", "Taux_Paye"), "R": ("M", "Taux_Jur"), "S": ("N", "Taux_PA")}
for ligne in range(2, NB_LIGNES + 2):
    ex = exemples[ligne - 2] if ligne - 2 < len(exemples) else None
    if ex:
        saisies = ex[:15] + [None] * 8 + ex[15:]
        for col, v in enumerate(saisies, start=1):
            if v is not None and v != "" and colonnes[col - 1][2] == "saisie":
                wc.cell(ligne, col, v)
    for col, (nouv, (act, taux)) in zip("PQRS", TAUX.items()):
        wc[f"{col}{ligne}"] = (
            f'=IF(OR($A{ligne}="",N({act}{ligne})=0),0,'
            f'CEILING({act}{ligne}*(1+IF($O{ligne}="",{taux},$O{ligne}))-0.005,Pas_Arrondi))')
    wc[f"T{ligne}"] = f'=IF($A{ligne}="","",SUM(K{ligne}:N{ligne}))'
    wc[f"U{ligne}"] = f'=IF($A{ligne}="","",SUM(P{ligne}:S{ligne}))'
    wc[f"V{ligne}"] = f'=IF($A{ligne}="","",U{ligne}-T{ligne})'
    wc[f"W{ligne}"] = f'=IF(OR($A{ligne}="",N(T{ligne})=0),"",V{ligne}/T{ligne})'
    # Ordre de passage dans l'onglet « Impression en masse »
    wc[f"Z{ligne}"] = (f'=IF(AND($A{ligne}<>"",$X{ligne}="Oui",OR(Filtre_Masse="Tous les clients à envoyer",'
                       f'$Y{ligne}<>"Email")),MAX(Z$1:Z{ligne - 1})+1,"")')
    for col in range(1, len(colonnes) + 1):
        c = wc.cell(ligne, col)
        typ = colonnes[col - 1][2]
        c.font = f(10, color=BLEU_SAISIE if typ == "saisie" else "000000")
        c.border = BORD
        if typ == "calc":
            c.fill = GRIS_CLAIR
        lettre = get_column_letter(col)
        if lettre in "KLMNPQRSTUV" and len(lettre) == 1:
            c.number_format = EUR
        elif lettre in ("O", "W"):
            c.number_format = PCT
        elif lettre == "H":
            c.number_format = "@"

dv_civ = DataValidation(type="list", formula1='"M.,Mme,M. et Mme"', allow_blank=True)
dv_oui = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
dv_mode = DataValidation(type="list", formula1='"Courrier,Email,Les deux"', allow_blank=True)
for dv, plage in ((dv_civ, "B"), (dv_oui, "X"), (dv_mode, "Y")):
    wc.add_data_validation(dv)
    dv.add(f"{plage}2:{plage}{NB_LIGNES + 1}")
wc.auto_filter.ref = f"A1:Z{NB_LIGNES + 1}"
wc["Z1"].comment = Comment("Numéro de page dans l'onglet « Impression en masse ». Vide = client non imprimé "
                           "(À envoyer ≠ Oui, ou mode Email si le filtre est « Courrier postal uniquement »).",
                           "Modèle")
wc["O1"].comment = Comment("Laisser vide pour appliquer les taux par défaut de l'onglet Paramètres. "
                           "Sinon, ce taux remplace tous les taux pour ce client (ex. 2,5 %).", "Modèle")
wc["N1"].comment = Comment("Honoraires annuels HT de la plateforme agréée (facturation électronique). "
                           "Laisser à 0 si le client n'y a pas souscrit.", "Modèle")
wb.defined_names["Codes_Clients"] = DefinedName("Codes_Clients",
                                                attr_text=f"Clients!$A$2:$A${NB_LIGNES + 1}")

# ---------------------------------------------------------------- Courrier (modèle réutilisable)
LIGNES_COURRIER = 45          # hauteur d'un courrier en lignes (= 1 page A4)
LARGEUR_CAR = 118             # caractères par ligne sur la largeur B:F
LARGEURS = {"A": 1.5, "B": 44, "C": 14, "D": 9, "E": 14, "F": 13, "G": 3,
            "H": 3, "I": 22, "J": 16, "K": 12, "L": 12, "M": 6, "N": 6, "O": 6}


def construire_courrier(ws, b, formule_ligne, formule_ref):
    """Dessine un courrier complet sur les lignes b+1 à b+45 de `ws`.

    formule_ligne : formule donnant la ligne du client dans l'onglet Clients (0 = aucun)
    formule_ref   : formule affichant la référence client en haut à droite
    Les colonnes J à O contiennent les calculs intermédiaires (non imprimés).
    """
    def R(n):
        return n + b

    L = f"$K${R(1)}"  # n° de ligne du client dans la plage Clients!x2:x501

    def cli(col):
        return f'IF({L}=0,"",INDEX(Clients!${col}$2:${col}$501,{L})&"")'

    def cli_num(col):
        return f"IF({L}=0,0,N(INDEX(Clients!${col}$2:${col}$501,{L})))"

    def balises(nom):
        return (f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({nom},'
                f'"{{cabinet}}",Cab_Nom),"{{societe}}",{cli("E")}),"{{date_effet}}",$K${R(3)}),'
                f'"{{hausse}}",$K${R(4)}),"{{civilite}}",$K${R(2)})')

    def paragraphe(ligne, nom):
        c = ws[f"B{R(ligne)}"]
        c.value = "=" + balises(nom)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="justify")
        ws.merge_cells(f"B{R(ligne)}:F{R(ligne)}")
        texte = next(p[2] for p in params if p[0] == nom)
        nb = max(1, math.ceil(len(texte) / LARGEUR_CAR))
        ws.row_dimensions[R(ligne)].height = 13 * nb + 11

    # Calculs intermédiaires
    ws[f"J{R(1)}"] = "Ligne client"
    ws[f"K{R(1)}"] = "=" + formule_ligne
    ws[f"J{R(2)}"] = "Civilité longue"
    ws[f"K{R(2)}"] = (f'=IF({cli("B")}="Mme","Madame",IF({cli("B")}="M.","Monsieur","Madame, Monsieur"))')
    ws[f"J{R(3)}"] = "Date effet texte"
    ws[f"K{R(3)}"] = ('=DAY(Date_Effet)&IF(DAY(Date_Effet)=1,"er","")&" "&INDEX(Liste_Mois,MONTH(Date_Effet))'
                      '&" "&YEAR(Date_Effet)')
    ws[f"J{R(4)}"] = "Hausse globale"
    ws[f"K{R(4)}"] = (f'=IF(SUM($K${R(28)}:$K${R(31)})=0,"",FIXED((SUM($L${R(28)}:$L${R(31)})'
                      f'/SUM($K${R(28)}:$K${R(31)})-1)*100,1)&" %")')

    # En-tête cabinet
    ws[f"B{R(2)}"] = "=Cab_Nom"
    ws[f"B{R(2)}"].font = f(15, True, BLEU_FONCE)
    for n, nom in ((3, "Cab_Adr1"), (4, "Cab_CPVille"), (5, "Cab_Contact")):
        ws[f"B{R(n)}"] = f"={nom}"
        ws[f"B{R(n)}"].font = f(9, color="404040")
    ws.row_dimensions[R(2)].height = 22
    for n in range(2, 6):
        ws.merge_cells(f"B{R(n)}:C{R(n)}")
    ws[f"E{R(2)}"] = "=" + formule_ref
    ws[f"E{R(2)}"].font = f(9, color="404040")
    ws[f"E{R(2)}"].alignment = Alignment(horizontal="right")
    ws.merge_cells(f"E{R(2)}:F{R(2)}")
    for col in "BCDEF":
        ws[f"{col}{R(6)}"].border = Border(bottom=Side(style="medium", color=BLEU_FONCE))
    ws.row_dimensions[R(6)].height = 6

    # Destinataire (fenêtre d'enveloppe à droite)
    for n in (7, 8, 9):
        ws.row_dimensions[R(n)].height = 14
    dest = {
        10: f'={cli("E")}',
        11: f'=IF({cli("D")}="","","À l\'attention de "&TRIM({cli("B")}&" "&{cli("C")}&" "&{cli("D")}))',
        12: f'={cli("F")}',
        13: f'=IF({cli("G")}="",TRIM({cli("H")}&" "&UPPER({cli("I")})),{cli("G")})',
        14: f'=IF({cli("G")}="","",TRIM({cli("H")}&" "&UPPER({cli("I")})))',
    }
    for n, formule in dest.items():
        ws[f"D{R(n)}"] = formule
        ws[f"D{R(n)}"].font = f(10.5, bold=(n == 10))
        ws.merge_cells(f"D{R(n)}:F{R(n)}")
        ws.row_dimensions[R(n)].height = 15
    ws.row_dimensions[R(15)].height = 26

    ws[f"D{R(16)}"] = ('=Cab_Ville&", le "&DAY(Date_Courrier)&IF(DAY(Date_Courrier)=1,"er","")&" "'
                       '&INDEX(Liste_Mois,MONTH(Date_Courrier))&" "&YEAR(Date_Courrier)')
    ws.merge_cells(f"D{R(16)}:F{R(16)}")
    ws.row_dimensions[R(17)].height = 20

    ws[f"B{R(18)}"] = "=" + balises("Txt_Objet")
    ws[f"B{R(18)}"].font = f(10.5, True)
    ws.merge_cells(f"B{R(18)}:F{R(18)}")
    ws.row_dimensions[R(19)].height = 14
    ws[f"B{R(20)}"] = f'=$K${R(2)}&","'
    ws.row_dimensions[R(21)].height = 6

    for n, nom in ((22, "Txt_P1"), (23, "Txt_P2"), (24, "Txt_P3"), (25, "Txt_P4")):
        paragraphe(n, nom)
    ws.row_dimensions[R(26)].height = 4

    # Tableau des honoraires
    for col, t in {"B": "Mission", "C": "Actuel HT", "D": "Évol.", "E": "Nouveau HT", "F": "Écart HT"}.items():
        c = ws[f"{col}{R(27)}"]
        c.value = t
        c.font = f(9.5, True, "FFFFFF")
        c.fill = ENTETE
        c.border = BORD
        c.alignment = Alignment(horizontal="left" if col == "B" else "center", vertical="center")
    ws.row_dimensions[R(27)].height = 18

    missions = [("Lib_Compta", "K", "P"), ("Lib_Paye", "L", "Q"), ("Lib_Jur", "M", "R"), ("Lib_PA", "N", "S")]
    for col, t in zip("JKLMNO", ("Libellé", "Actuel", "Nouveau", "Actif", "Rang", "k")):
        ws[f"{col}{R(27)}"] = t
    plage = lambda col: f"${col}${R(28)}:${col}${R(31)}"  # noqa: E731
    for i, (lib, act, nouv) in enumerate(missions):
        rr = R(28 + i)
        ws[f"J{rr}"] = f"={lib}"
        ws[f"K{rr}"] = "=" + cli_num(act)
        ws[f"L{rr}"] = "=" + cli_num(nouv)
        ws[f"M{rr}"] = f"=IF(OR(K{rr}<>0,L{rr}<>0),1,0)"
        ws[f"N{rr}"] = f'=IF(M{rr}=1,SUM($M${R(28)}:M{rr}),"")'
        ws[f"O{rr}"] = i + 1
        # k-ième mission souscrite : les missions à 0 ne s'affichent pas
        idx = f"MATCH($O{rr},{plage('N')},0)"
        ws[f"B{rr}"] = f'=IFERROR(INDEX({plage("J")},{idx}),"")'
        ws[f"C{rr}"] = f'=IFERROR(INDEX({plage("K")},{idx}),"")'
        ws[f"E{rr}"] = f'=IFERROR(INDEX({plage("L")},{idx}),"")'
        ws[f"F{rr}"] = f'=IF(E{rr}="","",E{rr}-C{rr})'
        ws[f"D{rr}"] = f'=IF(OR(E{rr}="",N(C{rr})=0),"",E{rr}/C{rr}-1)'
        for col in "BCDEF":
            ws[f"{col}{rr}"].font = f(9.5)
            ws[f"{col}{rr}"].alignment = Alignment(horizontal="left" if col == "B" else "right",
                                                   vertical="center", wrap_text=(col == "B"))
        for col in "CEF":
            ws[f"{col}{rr}"].number_format = EUR
        ws[f"D{rr}"].number_format = '+0.0%;-0.0%;"—"'
        ws.row_dimensions[rr].height = 17
    ws.conditional_formatting.add(f"B{R(28)}:F{R(31)}",
                                  FormulaRule(formula=[f'$B{R(28)}<>""'], border=Border(bottom=fin)))

    rt = R(32)
    ws[f"B{rt}"] = "Total annuel HT"
    ws[f"C{rt}"] = f"=SUM(K{R(28)}:K{R(31)})"
    ws[f"E{rt}"] = f"=SUM(L{R(28)}:L{R(31)})"
    ws[f"F{rt}"] = f"=E{rt}-C{rt}"
    ws[f"D{rt}"] = f'=IF(N(C{rt})=0,"",E{rt}/C{rt}-1)'
    for col in "BCDEF":
        c = ws[f"{col}{rt}"]
        c.font = f(10, True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.border = Border(top=Side(style="thin", color=BLEU_FONCE), bottom=Side(style="thin", color=BLEU_FONCE))
        c.alignment = Alignment(horizontal="left" if col == "B" else "right", vertical="center")
    for col in "CEF":
        ws[f"{col}{rt}"].number_format = EUR
    ws[f"D{rt}"].number_format = '+0.0%;-0.0%;"—"'
    ws.row_dimensions[rt].height = 18

    ws[f"B{R(33)}"] = (f'=IF(N(Nb_Mensualites)=0,"","Soit "&FIXED(E{rt}/Nb_Mensualites,2)&" € HT par échéance ("'
                       f'&Nb_Mensualites&" échéances), "&FIXED(E{rt}/Nb_Mensualites*(1+Taux_TVA),2)&" € TTC.")')
    ws[f"B{R(33)}"].font = f(9, italic=True, color="404040")
    ws.merge_cells(f"B{R(33)}:F{R(33)}")
    ws.row_dimensions[R(34)].height = 8

    for n, nom in ((35, "Txt_P5"), (36, "Txt_P6"), (37, "Txt_Politesse")):
        paragraphe(n, nom)

    ws[f"D{R(39)}"] = "=Signataire"
    ws[f"D{R(39)}"].font = f(10.5, True)
    ws[f"D{R(40)}"] = "=Signataire_Titre"
    ws[f"D{R(40)}"].font = f(9.5, italic=True)
    ws.merge_cells(f"D{R(39)}:F{R(39)}")
    ws.merge_cells(f"D{R(40)}:F{R(40)}")
    for n in (41, 42, 43, 44):
        ws.row_dimensions[R(n)].height = 16

    ws[f"B{R(45)}"] = "=Cab_Mentions"
    ws[f"B{R(45)}"].font = f(7.5, color="7F7F7F")
    ws[f"B{R(45)}"].alignment = Alignment(horizontal="center", wrap_text=True, vertical="bottom")
    ws.merge_cells(f"B{R(45)}:F{R(45)}")
    ws.row_dimensions[R(45)].height = 24
    for col in "BCDEF":
        ws[f"{col}{R(45)}"].border = Border(top=Side(style="thin", color="BFBFBF"))

    # Polices par défaut
    for row in ws.iter_rows(min_row=R(1), max_row=R(45), min_col=2, max_col=6):
        for c in row:
            if c.font.name != FONT:
                c.font = f(10)
    for row in ws.iter_rows(min_row=R(1), max_row=R(31), min_col=10, max_col=15):
        for c in row:
            c.font = f(8, color="A6A6A6")


def mise_en_page_a4(ws):
    ws.sheet_view.showGridLines = False
    for col, larg in LARGEURS.items():
        ws.column_dimensions[col].width = larg
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.6
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.page_margins.header = ws.page_margins.footer = 0.2


# ---- Onglet « Courrier » : un client choisi dans la liste
wl = wb.create_sheet("Courrier")
mise_en_page_a4(wl)
wl["I2"] = "CLIENT À ÉDITER ▼"
wl["I2"].font = f(10, True, "C00000")
wl["I3"] = "C001"
wl["I3"].font = f(12, True, BLEU_SAISIE)
wl["I3"].fill = JAUNE
wl["I3"].border = BORD
wl["I3"].alignment = Alignment(horizontal="center")
dv_cli = DataValidation(type="list", formula1="=Codes_Clients", allow_blank=False)
wl.add_data_validation(dv_cli)
dv_cli.add("I3")
wb.defined_names["Client_Selectionne"] = DefinedName("Client_Selectionne", attr_text="Courrier!$I$3")
wl["I4"] = '=IFERROR(INDEX(Clients!$E$2:$E$501,MATCH(Client_Selectionne,Codes_Clients,0)),"Code introuvable")'
wl["I6"] = ("Choisissez le code client dans la liste, puis :\n"
            "Fichier > Exporter > PDF (1 client)\nou Fichier > Imprimer.\n"
            "Pour tous les clients : onglet « Impression en masse ».")
construire_courrier(wl, 0, "IFERROR(MATCH(Client_Selectionne,Codes_Clients,0),0)",
                    '"Réf. client : "&Client_Selectionne')
wl["I4"].font = f(9, italic=True)
wl["I6"].font = f(9, color="595959")
wl["I6"].alignment = Alignment(wrap_text=True, vertical="top")
wl.merge_cells("I6:L10")
wl.print_area = f"A1:G{LIGNES_COURRIER}"
wl.page_setup.fitToWidth = 1
wl.page_setup.fitToHeight = 1
wl.sheet_properties.pageSetUpPr.fitToPage = True

# ---- Onglet « Impression en masse » : tous les courriers à la suite, 1 par page
NB_MASSE = 300
wi = wb.create_sheet("Impression en masse")
mise_en_page_a4(wi)
wi.page_setup.fitToWidth = 1
wi.page_setup.fitToHeight = 0  # hauteur libre : sauts de page manuels tous les 45 lignes
wi.sheet_properties.pageSetUpPr.fitToPage = True
for k in range(1, NB_MASSE + 1):
    b = (k - 1) * LIGNES_COURRIER
    construire_courrier(
        wi, b,
        f"IFERROR(MATCH({k},Clients!$Z$2:$Z$501,0),0)",
        f'IF($K${b + 1}=0,"","Réf. client : "&INDEX(Clients!$A$2:$A$501,$K${b + 1}))')
    wi.row_breaks.append(Break(id=b + LIGNES_COURRIER))
# Zone d'impression dynamique : seulement les courriers remplis
# (Excel n'imprime ainsi que les pages des clients sélectionnés, pas les 300 blocs)
wb.defined_names["Nb_Courriers"] = DefinedName(
    "Nb_Courriers", attr_text=f"MAX(Clients!$Z$2:$Z${NB_LIGNES + 1})")
wi.defined_names["_xlnm.Print_Area"] = DefinedName(
    "_xlnm.Print_Area", localSheetId=wb.sheetnames.index("Impression en masse"),
    attr_text=f"OFFSET('Impression en masse'!$A$1,0,0,{LIGNES_COURRIER}*MAX(1,Nb_Courriers),7)")
wi.sheet_properties.tabColor = "C00000"
wi["I2"] = "COURRIERS À IMPRIMER"
wi["I2"].font = f(10, True, "C00000")
wi["I3"] = "=Nb_Courriers"
wi["I3"].font = f(14, True, BLEU_SAISIE)
wi["I3"].fill = JAUNE
wi["I3"].alignment = Alignment(horizontal="center")
wi["I5"] = ("Fichier > Imprimer : tous les courriers sortent d'un coup (1 page par client).\n"
            "Fichier > Exporter > PDF : un seul PDF avec tous les courriers.\n"
            "Si des pages vierges apparaissent, imprimez les pages 1 à N (N = nombre ci-dessus).\n"
            "Sélection des clients : colonnes « À envoyer » et « Mode d'envoi » de l'onglet Clients.")
wi["I5"].font = f(9, color="595959")
wi["I5"].alignment = Alignment(wrap_text=True, vertical="top")
wi.merge_cells("I5:L12")


# ---------------------------------------------------------------- Mode d'emploi
wm = wb.create_sheet("Mode d'emploi", 0)
wm.sheet_view.showGridLines = False
wm.column_dimensions["A"].width = 3
wm.column_dimensions["B"].width = 110
lignes = [
    ("COURRIER D'AUGMENTATION TARIFAIRE — PUBLIPOSTAGE", f(16, True, BLEU_FONCE)),
    ("Compta · Paye · Juridique · Plateforme agréée", f(11, italic=True, color="595959")),
    ("", None),
    ("1. Onglet « Paramètres »", f(11, True)),
    ("Renseignez les coordonnées du cabinet, le signataire, la date d'effet, les taux d'augmentation par mission "
     "et, si besoin, adaptez les textes du courrier. Les balises {cabinet}, {societe}, {date_effet}, {hausse} "
     "et {civilite} sont remplacées automatiquement.", f(10)),
    ("", None),
    ("2. Onglet « Clients »", f(11, True)),
    ("Une ligne par client. Colonnes bleues = saisie ; colonnes vertes = calcul automatique (ne pas modifier). "
     "Saisissez les honoraires ANNUELS HT actuels par mission (0 ou vide si la mission n'est pas souscrite : "
     "elle n'apparaîtra pas dans le courrier). « Taux perso » remplace les taux par défaut pour ce client. "
     "Les 3 lignes d'exemple (C001 à C003) sont à remplacer par vos clients.", f(10)),
    ("", None),
    ("3. Onglet « Courrier » — un client à la fois", f(11, True)),
    ("Choisissez le code client dans la cellule jaune I3 : le courrier se remplit tout seul. "
     "Puis Fichier > Exporter > Créer un PDF (ou Enregistrer sous > PDF) pour l'envoi par e-mail, "
     "ou Fichier > Imprimer pour l'envoi postal. La mise en page A4 est déjà réglée (1 page, "
     "adresse du destinataire positionnée pour une enveloppe à fenêtre à droite).", f(10)),
    ("", None),
    ("4. IMPRESSION EN MASSE — onglet « Impression en masse » (sans macro)", f(11, True)),
    ("Cet onglet contient tous les courriers les uns à la suite des autres, un par page A4. Y figurent les clients "
     "dont « À envoyer » = Oui (et, par défaut, dont le mode d'envoi est Courrier ou Les deux — réglable dans "
     "Paramètres > Impression en masse). Cliquez sur l'onglet puis Fichier > Imprimer : toutes les lettres "
     "sortent d'un coup, recto simple. Pour un PDF unique : Fichier > Exporter > PDF avec cet onglet actif. "
     "La colonne Z de l'onglet Clients indique le numéro de page de chaque client. Capacité : 300 courriers.",
     f(10)),
    ("", None),
    ("5. Un PDF séparé par client — macro « Macro PDF » (facultatif)", f(11, True)),
    ("Enregistrez le fichier au format .xlsm (Classeur Excel prenant en charge les macros), ouvrez l'éditeur VBA "
     "(Alt + F11), Insertion > Module, collez le code de l'onglet « Macro PDF » (ou importez le fichier "
     "Module_Courriers.bas). Lancez ensuite (Alt + F8) :\n"
     "  • ExporterTousLesPDF : un PDF par client dont « À envoyer » = Oui, dans le dossier Courriers_PDF\n"
     "  • ExporterUnSeulPDF : un seul PDF regroupant tous les courriers (pratique pour un envoi en masse)\n"
     "  • ImprimerCourriersPostaux : impression des clients « À envoyer » = Oui et mode Courrier / Les deux", f(10)),
    ("", None),
    ("6. Alternative : publipostage Word", f(11, True)),
    ("L'onglet « Clients » (en-têtes en ligne 1) peut aussi servir de source de données à un publipostage "
     "Word (Publipostage > Sélection des destinataires > Utiliser une liste existante).", f(10)),
    ("", None),
    ("Légende des couleurs", f(11, True)),
    ("Texte bleu sur fond jaune = cellule à renseigner  ·  Fond gris = formule  ·  "
     "Colonnes grises à droite du courrier (J à O) = calculs internes, non imprimés.", f(10)),
]
for i, (txt, police) in enumerate(lignes, start=2):
    c = wm.cell(i, 2, txt)
    if police:
        c.font = police
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if len(txt) > 110 or "\n" in txt:
        wm.row_dimensions[i].height = 14 * (math.ceil(len(txt) / 120) + txt.count("\n")) + 4

# ---------------------------------------------------------------- Macro PDF
wv = wb.create_sheet("Macro PDF")
wv.column_dimensions["A"].width = 120
wv["A1"] = "Code VBA à coller dans un module (Alt + F11 > Insertion > Module), fichier enregistré en .xlsm"
wv["A1"].font = f(11, True, BLEU_FONCE)
with open("Module_Courriers.bas", encoding="utf-8") as fh:
    code = fh.read().splitlines()
for i, l in enumerate(code[1:], start=3):  # saute la ligne Attribute VB_Name
    c = wv.cell(i, 1, l)
    c.font = Font(name="Consolas", size=9)
    c.number_format = "@"
    if l.startswith("="):
        c.value = "'" + l

wb.active = 3  # ouvre sur le Courrier
wb.save(SORTIE)
print("OK ->", SORTIE)
