"""Crée le modèle Excel d'import des clients (Modele_import_clients.xlsx)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BORDEAUX = "9D202E"
F = "Arial"
COLONNES = [
    ("Réf. client", 12, "C001", "Facultatif. Laissé vide, l'application numérote C001, C002…"),
    ("Civilité", 10, "M.", "M., Mme ou M. et Mme. Vide = « Madame, Monsieur »."),
    ("Prénom", 14, "Pierre", None),
    ("Nom", 16, "DURAND", "Nom du contact. Utilisé pour « À l'attention de »."),
    ("Société", 30, "SARL DURAND BÂTIMENT", "Obligatoire (ou le Nom)."),
    ("Adresse", 30, "5 avenue des Tilleuls", None),
    ("Complément d'adresse", 22, "", "Bâtiment, BP, lieu-dit… Facultatif."),
    ("Code postal", 11, "95120", None),
    ("Ville", 18, "Ermont", None),
    ("E-mail", 26, "p.durand@exemple.fr", None),
    ("Comptabilité HT", 15, 3726, "Nouveaux honoraires annuels HT. 0 ou vide = mission non souscrite (non imprimée)."),
    ("Paye HT", 13, 1872, None),
    ("Juridique HT", 13, 618, None),
    ("Plateforme agréée HT", 15, 240, None),
    ("Périodicité", 14, "Mensuelle", "Mensuelle, Trimestrielle, Semestrielle ou Annuelle. Vide = périodicité par défaut du cabinet."),
    ("À envoyer", 11, "Oui", "Oui / Non. Vide = Oui."),
    ("Mode d'envoi", 13, "Courrier", "Courrier, Email ou Les deux. Vide = Courrier."),
]
fin = Side(style="thin", color="D9D9D9")

wb = Workbook()
ws = wb.active
ws.title = "Clients"
for i, (titre, larg, ex, com) in enumerate(COLONNES, start=1):
    c = ws.cell(1, i, titre)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BORDEAUX)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if com:
        c.comment = Comment(com, "ALTHO EXPERTISE")
    ws.column_dimensions[c.column_letter].width = larg
    e = ws.cell(2, i, ex if ex != "" else None)
    e.font = Font(name=F, color="0000FF", italic=True)
ws.row_dimensions[1].height = 30
for r in range(2, 502):
    for i in range(1, len(COLONNES) + 1):
        c = ws.cell(r, i)
        c.border = Border(bottom=fin)
        if r > 2:
            c.font = Font(name=F)
        if 11 <= i <= 14:
            c.number_format = '#,##0.00" €";-#,##0.00" €";""'
        if i == 8:
            c.number_format = "@"
ws["H2"].value = "95120"
ws.freeze_panes = "F2"
for formule, col in (('"M.,Mme,M. et Mme"', "B"), ('"Mensuelle,Trimestrielle,Semestrielle,Annuelle"', "O"), ('"Oui,Non"', "P"), ('"Courrier,Email,Les deux"', "Q")):
    dv = DataValidation(type="list", formula1=formule, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}501")

aide = wb.create_sheet("Mode d'emploi")
aide.column_dimensions["A"].width = 100
lignes = [
    ("Modèle d'import des clients — ALTHO EXPERTISE", Font(name=F, bold=True, size=14, color=BORDEAUX)),
    ("", None),
    ("1. Remplissez l'onglet « Clients » : une ligne par client, à partir de la ligne 2.", Font(name=F)),
    ("2. Remplacez ou supprimez la ligne d'exemple (en bleu).", Font(name=F)),
    ("3. Saisissez les NOUVEAUX honoraires annuels HT par mission (ce sont eux qui s'impriment sur la lettre).", Font(name=F)),
    ("   Laissez 0 ou vide une mission que le client n'a pas : elle n'apparaîtra pas dans son courrier.", Font(name=F)),
    ("4. Ne changez pas les titres de la ligne 1 (l'application s'en sert pour reconnaître les colonnes).", Font(name=F)),
    ("5. Enregistrez, puis dans l'application : onglet Clients > « Importer un fichier Excel ou CSV ».", Font(name=F)),
    ("", None),
    ("Colonnes obligatoires : Société (ou Nom), Adresse, Code postal, Ville. Les autres sont facultatives.", Font(name=F, italic=True, color="595959")),
]
for i, (t, f) in enumerate(lignes, start=1):
    aide.cell(i, 1, t)
    if f:
        aide.cell(i, 1).font = f

sortie = Path(__file__).parent.parent / "Modele_import_clients.xlsx"
wb.save(sortie)
print(sortie.name)
