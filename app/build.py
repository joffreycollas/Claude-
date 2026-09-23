"""Assemble l'application en un seul fichier HTML autonome (polices embarquées)."""
import base64
from pathlib import Path

ICI = Path(__file__).parent
POLICES = {
    "/*FONT_SANS*/": "HankenGrotesk-normal-latin.woff2",
    "/*FONT_SERIF*/": "SourceSerif4-normal-latin.woff2",
    "/*FONT_SERIF_ITALIC*/": "SourceSerif4-italic-latin.woff2",
}
IMAGES = {
    "/*IMG_ENTETE*/": ICI / "assets" / "entete.png",
    "/*IMG_PIED*/": ICI / "assets" / "pied.png",
    "/*IMG_LOGO*/": ICI / "assets" / "logo.png",
    "/*MODELE_XLSX*/": ICI.parent / "Modele_import_clients.xlsx",  # produit par modele_import.py
}
b64 = lambda chemin: base64.b64encode(chemin.read_bytes()).decode()
out = (ICI / "index.src.html").read_text(encoding="utf-8")
for marque, fichier in POLICES.items():
    out = out.replace(marque, b64(ICI / "fonts" / fichier))
for marque, chemin in IMAGES.items():
    assert marque in out, marque
    out = out.replace(marque, b64(chemin))
# Version personnelle : python build.py --clients liste.xlsx  ->  Courriers_ALTHO_avec_clients.html (non versé)
import hashlib, json, sys
embed, dest = "null", ICI.parent / "Courriers_ALTHO.html"
if "--clients" in sys.argv:
    from openpyxl import load_workbook
    src = Path(sys.argv[sys.argv.index("--clients") + 1])
    ws = load_workbook(src, data_only=True).worksheets[0]
    cell = lambda v: "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
    rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    empreinte = hashlib.sha1(json.dumps(rows).encode()).hexdigest()[:8]
    embed = json.dumps({"id": f"{src.name}:{len(rows)}:{empreinte}", "name": src.stem, "rows": rows}, ensure_ascii=False)
    dest = ICI.parent / "Courriers_ALTHO_avec_clients.html"
out = out.replace("/*EMBED_CLIENTS*/null", embed.replace("</", "<\\/"))
dest.write_text(out, encoding="utf-8")
print(f"{dest.name} : {len(out) // 1024} Ko")
